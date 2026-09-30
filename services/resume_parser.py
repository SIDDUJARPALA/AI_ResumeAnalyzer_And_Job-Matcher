from io import BytesIO
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError


class ResumeFileError(ValueError):
    pass


def extract_resume_text(file_name, content, config):
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension not in config["ALLOWED_EXTENSIONS"]:
        raise ResumeFileError("Upload a resume in PDF or DOCX format.")
    if not content:
        raise ResumeFileError("The uploaded resume is empty.")

    if extension == "pdf":
        text = _extract_pdf(content, config["MAX_PDF_PAGES"])
    else:
        text = _extract_docx(content, config["MAX_DOCX_UNCOMPRESSED_BYTES"])

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise ResumeFileError(
            "No readable text was found. Scanned image-only resumes are not supported."
        )
    if len(text) > config["MAX_RESUME_CHARS"]:
        raise ResumeFileError(
            f"The extracted resume is longer than the supported "
            f"{config['MAX_RESUME_CHARS']:,} characters."
        )
    return text


def _extract_pdf(content, max_pages):
    if not content.startswith(b"%PDF-"):
        raise ResumeFileError("The uploaded file is not a valid PDF.")
    try:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            if document.is_encrypted:
                raise ResumeFileError("Password-protected PDFs are not supported.")
            if len(document) > max_pages:
                raise ResumeFileError(f"PDF resumes may contain at most {max_pages} pages.")
            return "\n".join(page.get_text() for page in document)
    except ResumeFileError:
        raise
    except (pymupdf.FileDataError, ValueError) as error:
        raise ResumeFileError("The PDF could not be read. Please upload a valid PDF.") from error


def _extract_docx(content, max_uncompressed_bytes):
    if not content.startswith(b"PK"):
        raise ResumeFileError("The uploaded file is not a valid DOCX document.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > max_uncompressed_bytes
                or "word/document.xml" not in archive.namelist()
            ):
                raise ResumeFileError("The DOCX file is invalid or exceeds the supported size.")
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                cell.text for row in table.rows for cell in row.cells
            )
        return "\n".join(paragraphs)
    except ResumeFileError:
        raise
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError, XMLSyntaxError) as error:
        raise ResumeFileError("The DOCX could not be read. Please upload a valid DOCX.") from error
