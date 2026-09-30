import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document

from app import create_app
from services.resume_parser import extract_resume_text


class ApplicationFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(
            dir=Path(__file__).resolve().parents[1]
        )
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE_PATH": f"{self.temp_dir.name}/test.sqlite3",
                "LLM_PROVIDER": "groq",
                "GROQ_API_KEY": "groq-test-key",
            }
        )
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_upload_results_json_refinement_and_pdf_download(self):
        homepage = self.client.get("/")
        self.assertEqual(homepage.status_code, 200)
        self.assertIn(b"Your saved reports will appear here", homepage.data)
        self.assertIn(b"Groq", homepage.data)
        self.assertIn(b"Job description or requirements", homepage.data)
        self.assertIn(b"Skill gap analysis", homepage.data)
        self.assertIn(b"prepare for interviews", homepage.data)

        with patch("services.analysis_service.generate_json", side_effect=_analysis_responses()):
            response = self.client.post(
                "/analyze",
                data={
                    "resume": (BytesIO(_sample_pdf()), "sample.pdf"),
                    "job_description": "Python developer with SQL and two years of experience.",
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 302)
        results = self.client.get(response.headers["Location"])
        self.assertEqual(results.status_code, 200)
        self.assertIn(b"Transparent match score", results.data)
        self.assertIn(b"Python", results.data)
        self.assertIn(b"Python Developer", results.data)
        self.assertEqual(results.data.count(b'class="panel-card score-card"'), 1)
        self.assertIn(b"Transparent match score", results.data)
        self.assertIn(b"Resume improvement recommendations", results.data)
        self.assertIn(b"INTERVIEW PREPARATION", results.data)
        self.assertIn(b"COVER LETTER GENERATION", results.data)

        analysis_id = response.headers["Location"].rsplit("/", 1)[-1]
        api_result = self.client.get(f"/api/analysis/{analysis_id}")
        self.assertEqual(api_result.status_code, 200)
        result_json = json.loads(api_result.data)
        self.assertEqual(result_json["match"]["score"], 92)

        with patch(
            "services.analysis_service.generate_json",
            return_value={"revised_suggestions": ["Add a concise project summary."]},
        ):
            refined = self.client.post(
                f"/api/analysis/{analysis_id}/refine",
                json={"feedback": "Make the advice concise."},
            )
        self.assertEqual(refined.status_code, 200)
        self.assertEqual(
            json.loads(refined.data)["refinement"]["revised_suggestions"],
            ["Add a concise project summary."],
        )

        with patch(
            "services.analysis_service.generate_json",
            return_value={"revised_suggestions": ["Clarify the project impact."]},
        ) as generate:
            refined_again = self.client.post(
                f"/api/analysis/{analysis_id}/refine",
                json={"feedback": "Focus on project impact."},
            )
        refined_result = json.loads(refined_again.data)
        self.assertEqual(refined_result["refinement"]["revised_suggestions"], ["Clarify the project impact."])
        self.assertEqual(len(refined_result["refinement_history"]), 2)
        self.assertIn("Add a concise project summary.", generate.call_args.args[1])

        results = self.client.get(response.headers["Location"])
        self.assertIn(b"View 2 refinement rounds", results.data)

        report = self.client.get(f"/analysis/{analysis_id}/report.pdf")
        self.assertEqual(report.status_code, 200)
        self.assertTrue(report.data.startswith(b"%PDF-"))
        self.assertEqual(report.mimetype, "application/pdf")

        deleted = self.client.post(f"/analysis/{analysis_id}/delete")
        self.assertEqual(deleted.status_code, 302)
        self.assertEqual(self.client.get(f"/api/analysis/{analysis_id}").status_code, 404)
        homepage = self.client.get("/")
        self.assertNotIn(b"Python Developer", homepage.data)

    def test_api_rejects_non_pdf_or_docx_upload(self):
        response = self.client.post(
            "/api/analyze",
            data={
                "resume": (BytesIO(b"plain text"), "resume.txt"),
                "job_description": "A job description",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("PDF or DOCX", json.loads(response.data)["error"])

    def test_homepage_and_missing_api_key_error(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        self.app.config["GROQ_API_KEY"] = ""
        self.app.config["LLM_PROVIDER"] = "groq"

        response = self.client.post(
            "/api/analyze",
            data={
                "resume": (BytesIO(_sample_pdf()), "sample.pdf"),
                "job_description": "Python developer role.",
            },
            content_type="multipart/form-data",
        )

        self.assertEqual(response.status_code, 502)
        self.assertIn("GROQ_API_KEY", json.loads(response.data)["error"])

    def test_refinement_api_rejects_non_string_feedback(self):
        from database.models import save_analysis

        analysis_id = "53065345-5d36-4bba-a9a9-27485a10baac"
        save_analysis(
            self.app.config["DATABASE_PATH"],
            analysis_id,
            "sample.pdf",
            {
                "id": analysis_id,
                "resume": {},
                "job": {},
                "match": {},
                "recommendations": {"resume_improvements": []},
                "refinement": None,
            },
        )

        response = self.client.post(
            f"/api/analysis/{analysis_id}/refine",
            json={"feedback": ["not", "text"]},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("Feedback must", json.loads(response.data)["error"])

    def test_docx_resume_text_is_extracted(self):
        document = Document()
        document.add_paragraph("Taylor Example")
        document.add_paragraph("Skills: Python and SQL")
        content = BytesIO()
        document.save(content)

        extracted = extract_resume_text(
            "resume.docx", content.getvalue(), self.app.config
        )

        self.assertIn("Taylor Example", extracted)
        self.assertIn("Python and SQL", extracted)


def _sample_pdf():
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Alex Example\nPython Developer\nSkills: Python, SQL\n"
        "Education: Bachelor of Science in Computer Science\n"
        "Experience: 2 years building Flask APIs.",
    )
    content = document.tobytes()
    document.close()
    return content


def _analysis_responses():
    return [
        {
            "name": "Alex Example",
            "headline": "Python Developer",
            "skills": ["Python", "SQL", "Flask"],
            "education": ["Bachelor of Science in Computer Science"],
            "experience_years": 2,
            "experience_summary": ["Built Flask APIs for two years."],
            "projects": [],
            "strengths": ["Python development"],
        },
        {
            "title": "Python Developer",
            "required_skills": ["Python", "SQL", "Flask"],
            "minimum_experience_years": 3,
            "required_education": "Bachelor's degree",
            "key_responsibilities": ["Build Python services."],
        },
        {
            "resume_improvements": ["Add relevant details to the project section."],
            "interview_questions": ["Describe a Python API you built."],
            "cover_letter": "Dear Hiring Manager,\nI am applying for the Python Developer role.",
        },
    ]


if __name__ == "__main__":
    unittest.main()
